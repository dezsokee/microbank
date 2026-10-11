"""Builds the baseline / Istio / Linkerd benchmark workbook from sampler CSVs + k6 exports.

Usage: python benchmark/build_xlsx.py <results-dir> [<results-dir>...] <output.xlsx>

Several results dirs can be merged (e.g. Istio+baseline from one run.sh invocation and
Linkerd from another). The "Módszertan" sheet notes at the bottom of this file describe
the run they were written for — update them when re-running under different conditions.
"""
import csv
import json
import os
import sys
from collections import defaultdict
from datetime import datetime
from statistics import mean

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RES_DIRS, OUT_XLSX = sys.argv[1:-1], sys.argv[-1]


def rows_of(name):
    for d in RES_DIRS:
        p = os.path.join(d, name)
        if os.path.exists(p):
            with open(p) as f:
                yield from csv.DictReader(f)


def k6_file(stack, scen):
    for d in RES_DIRS:
        p = os.path.join(d, f"k6-{stack}-{scen}.json")
        if os.path.exists(p):
            return p
    raise FileNotFoundError(f"k6-{stack}-{scen}.json not found in {RES_DIRS}")
STACKS = ["baseline", "istio", "linkerd"]
LOAD = ["normal", "peak", "burst"]
PHASES = ["idle"] + LOAD
BUCKET = 10
PROXIES = {"istio-proxy", "linkerd-proxy"}
CATS = ["app", "sidecars", "control-plane", "monitoring", "ingress", "kubernetes"]
CAT_HU = {"app": "App konténerek", "sidecars": "Sidecar proxyk", "control-plane": "Mesh control plane",
          "monitoring": "Monitoring (Prometheus+Grafana)", "ingress": "Ingress NGINX", "kubernetes": "Kubernetes rendszer",
          "node": "Teljes minikube node"}


def category(ns, pod, container):
    if ns in ("istio-system", "linkerd"):
        return "control-plane"
    if ns == "ingress-nginx":
        return "ingress"
    if ns == "student-research":
        if pod.startswith("monitoring-"):
            return "monitoring"
        return "sidecars" if container in PROXIES else "app"
    return "kubernetes"


def service(pod):
    return "-".join(pod.split("-")[1:-2]) if pod.startswith("microbank-") else None


def keep(phase):
    return phase.split("/")[0] in STACKS

# ---------------- data ----------------
series = defaultdict(list)
for r in rows_of("containers.csv"):
    if keep(r["phase"]):
        series[(r["namespace"], r["pod"], r["container"])].append(
            (float(r["ts"]), r["phase"], float(r["cpu_seconds_total"]), int(r["mem_bytes"])))

cpu = defaultdict(lambda: defaultdict(float))
mem = defaultdict(lambda: defaultdict(float))
ts_rows = []  # per-container time series (rates)
for (ns, pod, cont), pts in series.items():
    pts.sort()
    uniq = [p for i, p in enumerate(pts) if i == 0 or p[0] != pts[i - 1][0]]
    cat, svc = category(ns, pod, cont), service(pod)
    keys = [cat] + ([f"svc:{svc}:{'sidecar' if cont in PROXIES else 'app'}"] if svc and ns == "student-research" else [])
    # The kubelet's sample timestamps drift relative to the sampler tick, so one
    # container occasionally lands two samples in the same 10 s bucket. Average per
    # container per bucket first, then sum across containers — summing raw samples
    # would count that container twice in that bucket.
    per_bucket = defaultdict(list)
    for (t0, _, c0, _), (t1, ph, c1, m1) in zip(uniq, uniq[1:]):
        if c1 < c0 or t1 <= t0:
            continue
        rate = (c1 - c0) / (t1 - t0)
        per_bucket[int(t1 // BUCKET)].append((ph, rate, m1))
        ts_rows.append((datetime.fromtimestamp(t1), ph.split("/")[0], ph.split("/")[1], ns, pod, cont, cat,
                        round(rate * 1000, 1), round(m1 / 2**20, 1)))
    for b, samples in per_bucket.items():
        ph = samples[-1][0]
        rate = mean(s[1] for s in samples)
        m1 = mean(s[2] for s in samples)
        for k in keys:
            cpu[(ph, b)][k] += rate
            mem[(ph, b)][k] += m1


def stats(store, ph, key):
    vals = [v.get(key, 0.0) for (p, _), v in store.items() if p == ph]
    return (mean(vals), max(vals)) if vals else (0.0, 0.0)


def seconds(ph):
    return len({b for (p, b) in cpu if p == ph}) * BUCKET


node = defaultdict(list)
node_rows = []
for r in rows_of("node.csv"):
    if keep(r["phase"]):
        c, m = float(r["cpu_percent"]) * 10, int(r["mem_bytes"]) / 2**20
        node[r["phase"]].append((c, m))
        st, ph = r["phase"].split("/")
        node_rows.append((datetime.fromtimestamp(float(r["ts"])), st, ph, round(c, 0), round(m, 0)))


def k6(stack, scen):
    m = json.load(open(k6_file(stack, scen)))["metrics"]
    d, ch = m["http_req_duration"], m.get("checks", {})
    return [m["http_reqs"]["rate"], m["http_reqs"]["count"], m["http_req_failed"]["value"],
            ch.get("passes", 0) / max(1, ch.get("passes", 0) + ch.get("fails", 0)),
            d["avg"], d["med"], d["p(90)"], d["p(95)"], d["p(99)"], d["max"],
            m["iteration_duration"]["avg"]]

# ---------------- styling helpers ----------------
F = "Arial"
HDR = PatternFill("solid", start_color="1F3864")
SUB = PatternFill("solid", start_color="D9E2F3")
INPUT = Font(name=F, color="0000FF")
BLACK = Font(name=F)
GREEN = Font(name=F, color="008000")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def header(ws, row, labels, col=1):
    for i, h in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=h)
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = HDR
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[row].height = 32


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(name=F, bold=True, size=14)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(name=F, italic=True, color="595959")


def cell(ws, r, c, v, fmt=None, font=None):
    x = ws.cell(row=r, column=c, value=v)
    x.font = font or BLACK
    x.border = BOX
    if fmt:
        x.number_format = fmt
    return x


def widths(ws, w):
    for i, v in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = v


wb = Workbook()

# ================= Sheet: Erőforrás (long table, inputs) =================
res = wb.active
res.title = "Erőforrás"
title(res, "Erőforrás-fogyasztás fázisonként és kategóriánként",
      "CPU: millicore (10 s-os vödrök átlaga / csúcsa). RAM: MiB working set. Forrás: kubelet Summary API + docker stats (mért adat, kék).")
header(res, 4, ["Stack", "Fázis", "Kategória", "Kategória (leírás)", "CPU átlag (m)", "CPU csúcs (m)",
                "RAM átlag (MiB)", "RAM csúcs (MiB)", "Fázis hossza (s)"])
r = 5
for st in STACKS:
    for ph in PHASES:
        key = f"{st}/{ph}"
        for c in CATS + ["node"]:
            if c == "node":
                n = node[key]
                ca, cp = mean(x for x, _ in n), max(x for x, _ in n)
                ma, mp = mean(y for _, y in n), max(y for _, y in n)
            else:
                ca, cp = (v * 1000 for v in stats(cpu, key, c))
                ma, mp = (v / 2**20 for v in stats(mem, key, c))
            vals = [st, ph, c, CAT_HU[c], round(ca, 1), round(cp, 1), round(ma, 1), round(mp, 1), seconds(key)]
            for i, v in enumerate(vals, 1):
                cell(res, r, i, v, "#,##0.0" if 5 <= i <= 8 else ("#,##0" if i == 9 else None), INPUT if i >= 5 else BLACK)
            r += 1
RES_LAST = r - 1
res.freeze_panes = "A5"
res.auto_filter.ref = f"A4:I{RES_LAST}"
widths(res, [11, 9, 15, 30, 13, 13, 14, 14, 13])

# ================= Sheet: Összehasonlítás (formulas over Erőforrás) =================
cmp_ = wb.create_sheet("Összehasonlítás")
title(cmp_, "Istio és Linkerd többlet a baseline-hoz képest (átlagos CPU és RAM)",
      "Minden érték képlet az 'Erőforrás' lapból (SUMIFS). Δ% = (mesh − baseline) / baseline; baseline 0 esetén '–'.")
header(cmp_, 4, ["Fázis", "Kategória",
                 "Baseline CPU (m)", "Istio CPU (m)", "Linkerd CPU (m)", "Istio Δ CPU (m)", "Linkerd Δ CPU (m)", "Istio Δ CPU %", "Linkerd Δ CPU %",
                 "Baseline RAM (MiB)", "Istio RAM (MiB)", "Linkerd RAM (MiB)", "Istio Δ RAM (MiB)", "Linkerd Δ RAM (MiB)", "Istio Δ RAM %", "Linkerd Δ RAM %"])
rng = lambda col: f"'Erőforrás'!${col}$5:${col}${RES_LAST}"
r = 5
for ph in PHASES:
    for c in CATS + ["node"]:
        cell(cmp_, r, 1, ph)
        cell(cmp_, r, 2, CAT_HU[c])
        cmp_.cell(row=r, column=17, value=c)  # helper key (hidden col Q)
        for base_col, metric in [(3, "E"), (10, "G")]:
            for off, stack in enumerate(STACKS):
                cell(cmp_, r, base_col + off,
                     f'=SUMIFS({rng(metric)},{rng("A")},"{stack}",{rng("B")},$A{r},{rng("C")},$Q{r})', "#,##0", GREEN)
            b, i, l = (get_column_letter(base_col + o) for o in range(3))
            cell(cmp_, r, base_col + 3, f"={i}{r}-{b}{r}", "#,##0;-#,##0;-")
            cell(cmp_, r, base_col + 4, f"={l}{r}-{b}{r}", "#,##0;-#,##0;-")
            cell(cmp_, r, base_col + 5, f'=IF({b}{r}=0,"–",{get_column_letter(base_col + 3)}{r}/{b}{r})', "0.0%;-0.0%;-")
            cell(cmp_, r, base_col + 6, f'=IF({b}{r}=0,"–",{get_column_letter(base_col + 4)}{r}/{b}{r})', "0.0%;-0.0%;-")
        if c == "node":
            for col in range(1, 17):
                x = cmp_.cell(row=r, column=col)
                x.fill = SUB
                x.font = Font(name=F, bold=True, color=x.font.color)
        r += 1
cmp_.column_dimensions["Q"].hidden = True
cmp_.freeze_panes = "C5"
widths(cmp_, [9, 30] + [12] * 14)

# ================= Sheet: k6 eredmények =================
k = wb.create_sheet("k6 eredmények", 0)
title(k, "MicroBank benchmark — baseline vs Istio vs Linkerd (k6 + erőforrás-hatékonyság)",
      "Kék: k6 summary export (mért). Fekete: képlet. Zöld: hivatkozás más lapra. CPU-ms/kérés = workload CPU (app + sidecar + control plane, monitoring/ingress/kube-system nélkül) × fázis hossza / kérésszám. Tisztán mesh-overhead = sidecar + control plane.")
cols = ["Forgatókönyv", "Stack", "Req/s", "Kérések", "Hibás kérés %", "Check siker %", "Átlag (ms)", "p50 (ms)",
        "p90 (ms)", "p95 (ms)", "p99 (ms)", "Max (ms)", "Iteráció átlag (ms)",
        "Workload CPU átlag: app + sidecar + control plane (m)", "Fázis hossza (s)", "CPU-ms / kérés"]
header(k, 4, cols)
r = 5
fmts = ["0.0", "#,##0", "0.00%", "0.0%", "#,##0", "#,##0", "#,##0", "#,##0", "#,##0", "#,##0", "#,##0"]
k6_rows = {}
for scen in LOAD:
    for st in STACKS:
        cell(k, r, 1, scen)
        cell(k, r, 2, st)
        for i, v in enumerate(k6(st, scen)):
            cell(k, r, 3 + i, v, fmts[i], INPUT)
        sum_terms = "+".join(
            f'SUMIFS({rng("E")},{rng("A")},$B{r},{rng("B")},$A{r},{rng("C")},"{c}")' for c in ("app", "sidecars", "control-plane"))
        cell(k, r, 14, f"={sum_terms}", "#,##0", GREEN)
        cell(k, r, 15, f'=SUMIFS({rng("I")},{rng("A")},$B{r},{rng("B")},$A{r},{rng("C")},"app")', "#,##0", GREEN)
        cell(k, r, 16, f"=IF(D{r}=0,0,N{r}*O{r}/D{r})", "0.0")
        k6_rows[(scen, st)] = r
        r += 1
K6_LAST = r - 1

r += 2
k.cell(row=r, column=1, value="Mesh overhead a baseline-hoz képest").font = Font(name=F, bold=True, size=12)
r += 1
header(k, r, ["Forgatókönyv", "Mesh", "Throughput Δ%", "p50 Δ%", "p95 Δ%", "p99 Δ%", "Hibaarány Δ (%-pont)",
              "Sidecar CPU átlag (m)", "Control plane CPU átlag (m)", "Sidecar RAM átlag (MiB)", "CPU-ms/kérés Δ%"])
OV_HDR = r
r += 1
for scen in LOAD:
    for mesh in ("istio", "linkerd"):
        b, i = k6_rows[(scen, "baseline")], k6_rows[(scen, mesh)]
        cell(k, r, 1, scen)
        cell(k, r, 2, mesh)
        for col, src in [(3, "C"), (4, "H"), (5, "J"), (6, "K")]:
            cell(k, r, col, f"=IF({src}{b}=0,0,{src}{i}/{src}{b}-1)", "0.0%;-0.0%;-")
        cell(k, r, 7, f"=(E{i}-E{b})*100", "0.00;-0.00;-")
        for col, metric, catg in [(8, "E", "sidecars"), (9, "E", "control-plane"), (10, "G", "sidecars")]:
            cell(k, r, col, f'=SUMIFS({rng(metric)},{rng("A")},$B{r},{rng("B")},$A{r},{rng("C")},"{catg}")', "#,##0", GREEN)
        cell(k, r, 11, f"=IF(P{b}=0,0,P{i}/P{b}-1)", "0.0%;-0.0%;-")
        r += 1
OV_LAST = r - 1
k.freeze_panes = "C5"
widths(k, [13, 10, 9, 10, 11, 11, 10, 9, 9, 9, 9, 9, 12, 18, 11, 12])

# chart: CPU-ms per request
ch = BarChart()
ch.type = "col"
ch.title = "CPU-ms / kérés (alacsonyabb = hatékonyabb)"
ch.y_axis.title = "CPU-ms"
data = Reference(k, min_col=16, min_row=4, max_row=K6_LAST)
ch.add_data(data, titles_from_data=True)
ch.set_categories(Reference(k, min_col=1, max_col=2, min_row=5, max_row=K6_LAST))
ch.height, ch.width = 8, 16
k.add_chart(ch, f"A{OV_LAST + 3}")

ch2 = BarChart()
ch2.type = "col"
ch2.title = "p95 latency (ms)"
ch2.add_data(Reference(k, min_col=10, min_row=4, max_row=K6_LAST), titles_from_data=True)
ch2.set_categories(Reference(k, min_col=1, max_col=2, min_row=5, max_row=K6_LAST))
ch2.height, ch2.width = 8, 16
k.add_chart(ch2, f"I{OV_LAST + 3}")

# ================= Sheet: Service bontás =================
sv = wb.create_sheet("Service bontás")
title(sv, "Service-enkénti átlagos CPU és RAM terhelés alatt", "app = alkalmazás konténer, sidecar = istio-proxy / linkerd-proxy (baseline-ban nincs).")
header(sv, 4, ["Stack", "Fázis", "Service", "Konténer", "CPU átlag (m)", "CPU csúcs (m)", "RAM átlag (MiB)", "RAM csúcs (MiB)"])
r = 5
svcs = sorted({kk.split(":")[1] for v in cpu.values() for kk in v if kk.startswith("svc:")})
for st in STACKS:
    for ph in PHASES:
        for s in svcs:
            for part in ("app", "sidecar"):
                key = f"svc:{s}:{part}"
                ca, cp = stats(cpu, f"{st}/{ph}", key)
                if part == "sidecar" and cp == 0:
                    continue
                ma, mp = stats(mem, f"{st}/{ph}", key)
                for i, v in enumerate([st, ph, s, part, round(ca * 1000, 1), round(cp * 1000, 1),
                                       round(ma / 2**20, 1), round(mp / 2**20, 1)], 1):
                    cell(sv, r, i, v, "#,##0.0" if i >= 5 else None, INPUT if i >= 5 else BLACK)
                r += 1
sv.freeze_panes = "A5"
sv.auto_filter.ref = f"A4:H{r - 1}"
widths(sv, [11, 9, 22, 10, 13, 13, 14, 14])

# ================= Sheet: Idősor - node =================
nd = wb.create_sheet("Idősor - node")
title(nd, "Teljes minikube node CPU/RAM idősor (docker stats, 10 s)")
header(nd, 4, ["Időpont", "Stack", "Fázis", "Node CPU (m)", "Node RAM (MiB)"])
for j, row in enumerate(sorted(node_rows), 5):
    for i, v in enumerate(row, 1):
        cell(nd, j, i, v, "yyyy-mm-dd hh:mm:ss" if i == 1 else ("#,##0" if i >= 4 else None), INPUT if i >= 4 else BLACK)
ND_LAST = 4 + len(node_rows)
nd.freeze_panes = "A5"
nd.auto_filter.ref = f"A4:E{ND_LAST}"
widths(nd, [20, 10, 11, 13, 15])
lc = LineChart()
lc.title = "Node CPU (m) — idősor (Istio, baseline, majd Linkerd szakaszok)"
lc.add_data(Reference(nd, min_col=4, min_row=4, max_row=ND_LAST), titles_from_data=True)
lc.height, lc.width = 9, 24
nd.add_chart(lc, "G4")

# ================= Sheet: Idősor - konténerek =================
ct = wb.create_sheet("Idősor - konténerek")
title(ct, "Konténerenkénti CPU-ráta és memória idősor (kubelet Summary API)")
header(ct, 4, ["Időpont", "Stack", "Fázis", "Namespace", "Pod", "Konténer", "Kategória", "CPU (m)", "RAM (MiB)"])
for j, row in enumerate(sorted(ts_rows), 5):
    for i, v in enumerate(row, 1):
        x = ct.cell(row=j, column=i, value=v)
        x.font = INPUT if i >= 8 else BLACK
        if i == 1:
            x.number_format = "yyyy-mm-dd hh:mm:ss"
        elif i >= 8:
            x.number_format = "#,##0.0"
ct.freeze_panes = "A5"
ct.auto_filter.ref = f"A4:I{4 + len(ts_rows)}"
widths(ct, [20, 10, 9, 17, 46, 22, 15, 10, 11])

# ================= Sheet: Módszertan =================
md = wb.create_sheet("Módszertan")
title(md, "Mérési módszertan és korlátok")
notes = [
    ("Környezet", "Egy node-os minikube (docker driver, 6 CPU, 10 GB RAM, 40 GB disk), Kubernetes v1.28.3, Windows 11 + WSL2 (12 GB). Istio 1.30.4 (Helm: base + istiod), STRICT PeerAuthentication + ISTIO_MUTUAL DestinationRule."),
    ("Baseline", "Ugyanaz a Helm release serviceMesh.provider=none-nal: nincs sidecar, istiod 0 replikára skálázva, mTLS policy-k törölve."),
    ("Linkerd", "Linkerd edge-26.5.3 a repóban lévő manifestből (k8s/linkerd/, ugyanazok a tanúsítványok), serviceMesh.provider=linkerd, namespace default-inbound-policy=all-authenticated + authorization-policies-strict.yaml (Server kivétel az api-gateway:8080 és frontend:80 ingress portokra). A manifest native sidecart használ, amit a K8s 1.28 nem támogat, ezért klasszikus sidecar módban fut (proxy.nativeSidecar=false) — ugyanúgy, mint az Istio sidecar. Az istiod 0 replikán."),
    ("Terhelés", "k6 (grafana/k6 Docker image) a minikube Docker hálózatról, közvetlenül az ingress-nginx-en (192.168.49.2). Szkriptek: k6/normal-traffic.js (10 VU, 5 p), k6/peak-traffic.js (50 VU, 9 p), k6/burst-traffic.js (100 VU, 4 p)."),
    ("Szkript-javítás", "Az auth-service felhasználónként egy tokent tárol; a VU-k korábban 3 seed useren osztoztak és kölcsönösen érvénytelenítették egymás tokenjét. Most minden VU dedikált usert kap (k6/lib/users.js, k6-vu-1..100, 1 000 000 EUR kezdőegyenleg)."),
    ("Futási sorrend", "Stackenként: konfigurálás + minden pod újraindítása (Postgres emptyDir → üres DB), 1 p warm-up (nem mért), 2 p idle, normal, 1 p szünet, peak, 1 p szünet, burst. Sorrend: Istio, baseline (egy menetben), majd Linkerd egy külön menetben, a gép újraindítása után (az első Linkerd próbálkozást memóriahiány szakította meg)."),
    ("Erőforrás-limitek", "A mérés idejére az account-service és transaction-service CPU limitje 500m → 1500m (mindkét stacknél azonos); minden más a Helm chart alapértéke. A kért döntés célja: a peak/burst ne a limitbe ütközzön."),
    ("Mintavétel", "10 s-onként: kubelet Summary API (konténerenkénti kumulatív CPU-nanoszekundum + working set), a CPU-ráta két minta különbségéből; node szint: docker stats minikube. A k6 konténer nincs benne a mérésben."),
    ("Kategóriák", "App: student-research meshelt podok alkalmazás konténerei + Postgres. Sidecar: istio-proxy / linkerd-proxy. Control plane: istio-system / linkerd namespace (Linkerd-nél a control plane podok saját proxyja is ide számít). Monitoring: Prometheus + Grafana (mesh alatt sidecarral). Kubernetes: kube-system."),
    ("Instrumentáció", "Az api-gateway és transaction-service a client_request_duration_seconds metrikával bővített lokális image-en fut mindkét stackben."),
    ("Ismert hibák", "Burst alatt a transaction-service readiness probe-ja (3 s timeout) időnként kiesik → 502 (baseline) / 503 (Istio). Ez valós telítettségi viselkedés, nem mérési hiba."),
    ("Korlátok", "Egyetlen futás kombinációnként (nincs ismétlés, szórás); a k6 ugyanazon a gépen fut; az otel collector nincs telepítve (exportálási hibák minden stackben jelen). A Linkerd menet más időpontban (újraindított gépen) futott, mint az Istio/baseline — ugyanaz a cluster és konfiguráció, de a host állapota eltér: a node-szintű (docker stats) CPU idle-ben a Linkerd menetben ~640m volt a ~200–270m helyett, miközben a konténerenkénti összeg ugyanolyan alacsony. A többlet a podokon kívüli folyamatokból jön (kubelet, dockerd a minikube konténerben), ezért a stackek közötti összevetéshez a konténer-szintű kategóriák (app, sidecar, control plane) a megbízhatóak, a 'Teljes minikube node' sor kevésbé."),
]
header(md, 3, ["Téma", "Leírás"])
for j, (a, b) in enumerate(notes, 4):
    cell(md, j, 1, a, font=Font(name=F, bold=True))
    x = cell(md, j, 2, b)
    x.alignment = Alignment(wrap_text=True, vertical="top")
    md.row_dimensions[j].height = 48
widths(md, [18, 120])

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
wb.save(OUT_XLSX)
print("saved", OUT_XLSX, "| erőforrás sorok:", RES_LAST - 4, "| idősor konténer sorok:", len(ts_rows))
